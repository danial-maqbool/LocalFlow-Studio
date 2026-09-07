"""Synthetic acceptance operations using each repository's own installed runtime."""

from pathlib import Path
import csv
import hashlib
import secrets
import io
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import time
import traceback
import zipfile

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))
from app.service import Application
from localdesk.safety import InputError
from tests.fixtures import make_image, make_pdf

import argparse

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "--report-dir",
    type=Path,
    required=True,
    help="New ignored evidence folder for this run. Existing folders are preserved.",
)
args = parser.parse_args()
BASE = args.report_dir.resolve()
BASE.mkdir(parents=True, exist_ok=False)
WORK = BASE / "synthetic space unicode-ÃƒÂ©"
WORK.mkdir()
RESULTS = []
APP = None
PASSWORD = secrets.token_urlsafe(32)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def record(name, func):
    started = time.monotonic()
    try:
        detail = func()
        result = dict(check=name, status="PASS", detail=detail)
    except Exception:
        result = dict(check=name, status="FAIL", traceback=traceback.format_exc())
    result["seconds"] = time.monotonic() - started
    RESULTS.append(result)
    (BASE / "results.json").write_text(json.dumps(RESULTS, indent=2, default=str), encoding="utf-8")
    print(name, result["status"], flush=True)


def finish(action, body, expected="done"):
    ident = APP.common_post(action, body)["job_id"]
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        job = APP.jobs.get(ident)
        if job["status"] not in ("running", "queued"):
            assert job["status"] == expected, job
            return job["result"] if expected == "done" else job
        time.sleep(0.03)
    raise TimeoutError(action)


def output(artifact):
    return APP.download(artifact["path"]).read_bytes()


def restart():
    global APP
    APP.close()
    APP = Application(
        ROOT, BASE / "persistent-data", passphrase=PASSWORD
    )


def file(name, content):
    path = WORK / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content.encode("utf-8") if isinstance(content, str) else content)
    return path


def office_fixtures():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    doc = Document()
    doc.add_paragraph("Office canary qa@example.test public number 42")
    doc.core_properties.author = "PRIVATE_AUTHOR_CANARY"
    doc.save(WORK / "office.docx")
    book = Workbook()
    book.active.append(["Office canary qa@example.test", 42])
    book.save(WORK / "office.xlsx")
    book.close()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = (
        "Office canary qa@example.test public number 42"
    )
    deck.save(WORK / "office.pptx")
    return [WORK / ("office" + ext) for ext in (".docx", ".xlsx", ".pptx")]


def localflow():
    def invoice():
        shutil.copytree(ROOT / "examples", WORK / "invoices")
        paths = list((WORK / "invoices").glob("invoice-*.txt"))
        before = {str(p): sha(p) for p in paths}
        preview = finish(
            "run", {"id": "invoice-register", "workspace": str(WORK / "invoices"), "dry_run": True}
        )
        assert not list(APP.exports.iterdir())
        assert not preview["artifacts"]
        run = finish(
            "run", {"id": "invoice-register", "workspace": str(WORK / "invoices"), "dry_run": False}
        )
        for item in run["artifacts"]:
            raw = output(item)
            if item["name"].endswith(".csv"):
                rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
                assert len(rows) == 3, rows
                expected = [
                    {
                        "invoice_id": p.read_text().splitlines()[0].split(": ", 1)[1],
                        "name": p.name,
                        "copied_as": p.read_text().splitlines()[0].split(": ", 1)[1] + ".txt",
                        "sha256": sha(p),
                    }
                    for p in sorted(paths)
                ]
                assert rows == expected, (rows, expected)
            elif item["name"].endswith(".txt"):
                assert raw in [p.read_bytes() for p in paths]
        assert before == {str(p): sha(p) for p in paths}
        return {"preview": preview, "run": run, "csv_rows": rows, "hashes": before}

    record("LF-01 copied fixture preview output CSV and copied bytes source hashes", invoice)

    def tables():
        from localdesk.documents import document_request
        from PIL import Image, ImageDraw, ImageFont

        ruled = make_pdf(WORK / "ruled.pdf", table=True)
        lines = document_request(ruled, "tables", strategy="lines")
        assert lines["tables"][0]["rows"] == [
            ["Item", "Total"],
            ["Invoice A", "42"],
            ["Invoice B", "75"],
        ], lines
        from pypdf import PdfReader, PdfWriter

        reader = PdfReader(str(ruled))
        page = reader.pages[0]
        stream = page.get_contents().get_data()
        # Remove only the ruled-line drawing prefix, keeping exact positioned text.
        from pypdf.generic import DecodedStreamObject, NameObject

        textstream = DecodedStreamObject()
        textstream.set_data(stream[stream.index(b"BT ") :])
        page[NameObject("/Contents")] = textstream
        writer = PdfWriter()
        writer.add_page(page)
        writer.write(WORK / "borderless.pdf")
        borderless = document_request(WORK / "borderless.pdf", "tables", strategy="text")
        found = [r for t in borderless["tables"] for r in t["rows"] if any(r)]
        assert any("42" in r for r in found) and any("75" in r for r in found), borderless
        im = Image.new("RGB", (1000, 400), "white")
        draw = ImageDraw.Draw(im)
        font = ImageFont.load_default(size=36)
        for y, left, right in [
            (40, "Item", "Total"),
            (140, "Invoice A", "42"),
            (240, "Invoice B", "75"),
        ]:
            draw.text((30, y), left, font=font, fill="black")
            draw.text((600, y), right, font=font, fill="black")
        im.save(WORK / "table-scan.pdf", "PDF", resolution=144)
        ocr = document_request(
            WORK / "table-scan.pdf", "tables", strategy="ocr", column_edges=[250]
        )
        rows = [r for t in ocr["tables"] for r in t["rows"]]
        assert ["Invoice A", "42"] in rows and ["Invoice B", "75"] in rows, ocr
        try:
            document_request(WORK / "table-scan.pdf", "text", ocr=True, language="zzz_missing")
        except InputError as exc:
            missing = str(exc)
        else:
            raise AssertionError("Missing OCR language was not rejected")
        return {
            "ruled": lines,
            "borderless": borderless,
            "ocr_columns": ocr,
            "missing_language": missing,
        }

    record("LF-03 ruled borderless configured OCR table values missing language", tables)

    def semantics():
        from localdesk.semantic import SemanticModel, extractive_summary

        corpus = [
            "Supplier invoice payment is due. Supplier receipt total is 42.",
            "Football players score goals. The football stadium has teams.",
            "Invoice supplier receipt payment totals. Payment is due today.",
        ]
        model = SemanticModel(corpus)
        scores = model.scores("supplier invoice payment")
        assert scores[0] > scores[1] and scores[2] > scores[1], scores
        assert model.scores("UNKNOWNQUERYXYZ") == [0.0, 0.0, 0.0]
        small = SemanticModel(corpus[:1])
        assert "TF-IDF" in small.method
        summary = extractive_summary(corpus[0], sentences=1)
        assert all(line in corpus[0] for line in summary.splitlines())
        return {
            "backend": model.method,
            "small_backend": small.method,
            "scores": scores,
            "summary": summary,
        }

    record("LF-04 explicit LSA TF-IDF unknown queries source sentence summary", semantics)


def shared():
    def persist():
        APP.store.set("qa_canary", "PERSISTED_SYNTHETIC_CANARY")
        backup = APP.common_post("backup", {})
        raw = output(backup)
        assert b"PERSISTED_SYNTHETIC_CANARY" not in raw
        restart()
        assert APP.store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        from localdesk.storage import Store

        restored = BASE / "restored.vault"
        restored.write_bytes(raw)
        store = Store(restored, PASSWORD)
        assert store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        store.close()
        before = sha(restored)
        try:
            Store(restored, secrets.token_urlsafe(32))
        except InputError:
            pass
        else:
            raise AssertionError("Wrong password accepted")
        assert sha(restored) == before
        return {"backup_sha256": hashlib.sha256(raw).hexdigest(), "restart_and_restore": True}

    record("Shared real encrypted persistence backup restore wrong-password preservation", persist)


def main():
    global APP
    APP = Application(
        ROOT, BASE / "persistent-data", passphrase=PASSWORD
    )
    try:
        localflow()
        shared()
    finally:
        APP.close()
        hashes = {
            p.relative_to(BASE).as_posix(): sha(p)
            for p in BASE.rglob("*")
            if p.is_file() and p.name != "hashes.json"
        }
        (BASE / "hashes.json").write_text(json.dumps(hashes, indent=2), encoding="utf-8")
    return int(any(r["status"] == "FAIL" for r in RESULTS))


if __name__ == "__main__":
    raise SystemExit(main())
